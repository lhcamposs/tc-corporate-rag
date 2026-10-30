"""
Script do experimento do TCC (Seção 3.9 da monografia).

Fluxo:
  1. python avaliar.py trechos     -> trechos.xlsx (todos os trechos da base, para anotar o gabarito)
  2. python avaliar.py modelo      -> gabarito.xlsx (planilha com as 60 perguntas, cotas por área e tipo)
     ... o autor escreve as perguntas, as respostas de referência e os vectorIds relevantes ...
  3. python avaliar.py validar     -> confere o gabarito antes de rodar o experimento
  4. python avaliar.py recuperar   -> Recall@1/3/5/10, MRR@10 e tempos de recuperação dos três métodos
  5. python avaliar.py gerar       -> respostas do pgvector, do ChromaDB e do teto + planilha de avaliação cega
     ... o autor dá as notas na planilha avaliacao_cega.xlsx ...
  6. python avaliar.py notas       -> nota média, degradação em relação ao teto, fundamentação e tempos de geração

Requer: Python 3.10+, openpyxl (pip install -r requirements.txt) e a aplicação rodando (padrão http://localhost:8080).
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import statistics
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("Instale as dependências: pip install -r requirements.txt")

AQUI = Path(__file__).resolve().parent
RESULTADOS = AQUI / "resultados"

METODOS_RECUPERACAO = ["pgvector", "chroma", "lexical"]
CONFIGS_GERACAO = ["pgvector", "chroma", "teto"]
KS = [1, 3, 5, 10]
K_MAX = 10

TIPOS = ["literal", "parafrase", "sem_resposta"]
NOTAS_VALIDAS = [1, 3, 5, 7, 10]
AREA_RELACIONAL = "Base relacional"
AREA_SEM_RESPOSTA = "Sem resposta"

# Áreas da Tabela 3.2 e cotas de perguntas com resposta (literal, paráfrase).
# Proporcionais ao número de trechos de cada área, com um mínimo para que todas apareçam.
COTAS = {
    "Gestão de pessoas": (7, 6),
    "Regimento interno": (5, 5),
    "Regulamentos acadêmicos": (4, 4),
    "Diárias e passagens": (4, 3),
    "Tecnologia da informação": (2, 3),
    "Instrução Normativa nº 1": (2, 1),
    AREA_RELACIONAL: (1, 3),
}
QTD_SEM_RESPOSTA = 10


def area_do_arquivo(arquivo: str) -> str:
    nome = (arquivo or "").lower()
    if nome == "artigo_conhecimento":
        return AREA_RELACIONAL
    if nome.startswith("regimento_interno"):
        return "Regimento interno"
    if nome.startswith(("regulamento_de_monitoria", "regulamento_dos_cursos_de_gradu")):
        return "Regulamentos acadêmicos"
    if nome.startswith(("manual_scdp", "tutorial_para_cancelamento")):
        return "Diárias e passagens"
    if nome.startswith(("psi", "norma_03", "norma_04", "campanha_de_conscientiza")):
        return "Tecnologia da informação"
    if nome.startswith("in_0001"):
        return "Instrução Normativa nº 1"
    return "Gestão de pessoas"


# ----------------------------------------------------------------------------- HTTP

class Api:
    def __init__(self, base: str, timeout: float):
        self.base = base.rstrip("/")
        self.timeout = timeout

    def _req(self, metodo: str, caminho: str, corpo: dict | None = None, params: dict | None = None):
        url = self.base + caminho
        if params:
            url += "?" + urllib.parse.urlencode(params)
        dados = json.dumps(corpo).encode("utf-8") if corpo is not None else None
        req = urllib.request.Request(url, data=dados, method=metodo,
                                     headers={"Content-Type": "application/json; charset=utf-8"})
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            detalhe = e.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"{metodo} {url} -> HTTP {e.code}: {detalhe}") from None
        except urllib.error.URLError as e:
            raise RuntimeError(f"Não foi possível acessar {url}: {e.reason}. A aplicação está rodando?") from None

    def trechos(self):
        return self._req("GET", "/api/evaluation/chunks")

    def recuperar(self, pergunta: str, metodo: str, k: int):
        return self._req("POST", "/api/evaluation/retrieve", {"pergunta": pergunta}, {"method": metodo, "k": k})

    def consultar(self, pergunta: str, store: str):
        return self._req("POST", "/api/query", {"pergunta": pergunta}, {"store": store})

    def teto(self, pergunta: str, vector_ids: list[str]):
        return self._req("POST", "/api/evaluation/ceiling", {"pergunta": pergunta, "vectorIds": vector_ids})


# ----------------------------------------------------------------------------- planilhas

CAB = Font(bold=True, color="FFFFFF")
CAB_FUNDO = PatternFill("solid", fgColor="2F5496")
QUEBRA = Alignment(wrap_text=True, vertical="top")


def cabecalho(ws, colunas: list[tuple[str, int]]):
    for i, (nome, largura) in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font, c.fill = CAB, CAB_FUNDO
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"


def cmd_trechos(args, api: Api):
    trechos = api.trechos()
    wb = Workbook()
    ws = wb.active
    ws.title = "Trechos"
    cabecalho(ws, [("vector_id", 38), ("area", 24), ("arquivo", 40), ("chunk_index", 8), ("texto", 110)])
    for t in trechos:
        ws.append([t["vectorId"], area_do_arquivo(t["arquivo"]), t["arquivo"], t["chunkIndex"], t["trecho"]])
        for c in ws[ws.max_row]:
            c.alignment = QUEBRA
    ws.auto_filter.ref = ws.dimensions
    destino = Path(args.saida)
    wb.save(destino)
    por_area = defaultdict(int)
    for t in trechos:
        por_area[area_do_arquivo(t["arquivo"])] += 1
    print(f"{len(trechos)} trechos exportados para {destino}")
    for area, n in sorted(por_area.items(), key=lambda x: -x[1]):
        print(f"  {area}: {n}")


COLUNAS_GABARITO = [("id", 7), ("tipo", 13), ("area", 24), ("pergunta", 60), ("resposta_referencia", 60),
                    ("trechos_relevantes", 42), ("observacoes", 30)]


def cmd_modelo(args, _api=None):
    destino = Path(args.saida)
    if destino.exists() and not args.sobrescrever:
        sys.exit(f"{destino} já existe. Use --sobrescrever para recriar (as perguntas escritas serão perdidas).")
    wb = Workbook()

    inst = wb.active
    inst.title = "Instrucoes"
    texto = [
        "Gabarito do experimento (Seção 3.9 da monografia)",
        "",
        "1. Na aba Perguntas, cada linha já tem o tipo e a área definidos pelas cotas. Escreva a pergunta e a resposta de referência.",
        "2. literal: usa as mesmas palavras do trecho que a responde. parafrase: a mesma dúvida com outras palavras (evite repetir os termos do trecho).",
        "3. sem_resposta: pergunta plausível sobre a instituição cuja resposta NÃO está na base. Deixe trechos_relevantes vazio.",
        "4. trechos_relevantes: os vector_id (copiados de trechos.xlsx) de TODOS os trechos que contêm a resposta, separados por ponto e vírgula.",
        "   Se a resposta estiver no registro da base relacional e no PDF de origem, marque os dois.",
        "5. resposta_referencia: curta e completa, como você responderia (ex.: '45 dias para docentes EBTT e 30 para técnicos').",
        "6. A aba Cotas mostra quantas perguntas de cada tipo e área já foram escritas.",
        "7. Termine o gabarito ANTES de rodar 'recuperar' e 'gerar', para que ele não seja influenciado pelos resultados.",
        "8. Rode 'python avaliar.py validar' para conferir tudo.",
    ]
    for linha in texto:
        inst.append([linha])
    inst["A1"].font = Font(bold=True, size=13)
    inst.column_dimensions["A"].width = 140

    ws = wb.create_sheet("Perguntas")
    cabecalho(ws, COLUNAS_GABARITO)
    linhas = []
    for area, (lit, par) in COTAS.items():
        linhas += [("literal", area)] * lit + [("parafrase", area)] * par
    linhas += [("sem_resposta", AREA_SEM_RESPOSTA)] * QTD_SEM_RESPOSTA
    for i, (tipo, area) in enumerate(linhas, start=1):
        ws.append([f"P{i:02d}", tipo, area, None, None, None, None])
        for c in ws[ws.max_row]:
            c.alignment = QUEBRA
    n = len(linhas) + 1
    dv_tipo = DataValidation(type="list", formula1='"' + ",".join(TIPOS) + '"', allow_blank=False)
    dv_area = DataValidation(type="list", formula1='"' + ",".join(list(COTAS) + [AREA_SEM_RESPOSTA]) + '"')
    ws.add_data_validation(dv_tipo)
    ws.add_data_validation(dv_area)
    dv_tipo.add(f"B2:B{n}")
    dv_area.add(f"C2:C{n}")
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    ws.conditional_formatting.add(f"D2:E{n}", FormulaRule(formula=["LEN(D2)=0"], fill=amarelo))
    ws.conditional_formatting.add(f"F2:F{n}", FormulaRule(formula=['AND(LEN(F2)=0,$B2<>"sem_resposta")'], fill=amarelo))

    cot = wb.create_sheet("Cotas")
    cabecalho(cot, [("area", 26), ("literal (meta)", 14), ("literal (escritas)", 16),
                    ("parafrase (meta)", 16), ("parafrase (escritas)", 18)])
    for area, (lit, par) in list(COTAS.items()) + [(AREA_SEM_RESPOSTA, (QTD_SEM_RESPOSTA, 0))]:
        r = cot.max_row + 1
        tipo_lit = "sem_resposta" if area == AREA_SEM_RESPOSTA else "literal"
        cot.append([area, lit,
                    f'=COUNTIFS(Perguntas!$C:$C,$A{r},Perguntas!$B:$B,"{tipo_lit}",Perguntas!$D:$D,"<>")',
                    par,
                    f'=COUNTIFS(Perguntas!$C:$C,$A{r},Perguntas!$B:$B,"parafrase",Perguntas!$D:$D,"<>")'])
    r = cot.max_row + 1
    cot.append(["Total", f"=SUM(B2:B{r-1})", f"=SUM(C2:C{r-1})", f"=SUM(D2:D{r-1})", f"=SUM(E2:E{r-1})"])
    for c in cot[r]:
        c.font = Font(bold=True)
    cot.cell(row=r + 2, column=1, value="Na linha 'Sem resposta', a coluna 'literal' conta as perguntas sem resposta.")

    wb.active = 1
    wb.save(destino)
    print(f"Planilha criada: {destino} ({len(linhas)} perguntas: "
          f"{sum(l for l, _ in COTAS.values())} literais, {sum(p for _, p in COTAS.values())} paráfrases, "
          f"{QTD_SEM_RESPOSTA} sem resposta)")


def ler_gabarito(caminho: Path, exigir_completo: bool = True) -> list[dict]:
    if not caminho.exists():
        sys.exit(f"Gabarito não encontrado: {caminho}. Rode 'python avaliar.py modelo' primeiro.")
    ws = load_workbook(caminho, data_only=True)["Perguntas"]
    nomes = [c.value for c in ws[1]]
    perguntas = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        reg = dict(zip(nomes, linha))
        if not reg.get("id"):
            continue
        texto = (reg.get("pergunta") or "").strip()
        if not texto and not exigir_completo:
            continue
        ids = [x.strip() for x in str(reg.get("trechos_relevantes") or "").replace(",", ";").replace("\n", ";").split(";")
               if x.strip()]
        perguntas.append({
            "id": str(reg["id"]).strip(),
            "tipo": (reg.get("tipo") or "").strip(),
            "area": (reg.get("area") or "").strip(),
            "pergunta": texto,
            "referencia": (reg.get("resposta_referencia") or "").strip(),
            "relevantes": ids,
        })
    return perguntas


def cmd_validar(args, api: Api | None):
    perguntas = ler_gabarito(Path(args.gabarito), exigir_completo=True)
    erros, avisos = [], []
    ids_validos = None
    if api is not None and not args.offline:
        try:
            ids_validos = {t["vectorId"]: t for t in api.trechos()}
        except RuntimeError as e:
            avisos.append(f"não foi possível conferir os vector_id na aplicação ({e}); use --offline para pular")
    vistos = set()
    for p in perguntas:
        pid = p["id"]
        if pid in vistos:
            erros.append(f"{pid}: id repetido")
        vistos.add(pid)
        if p["tipo"] not in TIPOS:
            erros.append(f"{pid}: tipo inválido '{p['tipo']}'")
        if not p["pergunta"]:
            erros.append(f"{pid}: pergunta ainda não escrita")
            continue
        if p["tipo"] != "sem_resposta":
            if not p["referencia"]:
                erros.append(f"{pid}: falta a resposta de referência")
            if not p["relevantes"]:
                erros.append(f"{pid}: falta trechos_relevantes")
        elif p["relevantes"]:
            erros.append(f"{pid}: pergunta sem_resposta não deve ter trechos_relevantes")
        if ids_validos is not None:
            for vid in p["relevantes"]:
                if vid not in ids_validos:
                    erros.append(f"{pid}: vector_id inexistente na base: {vid}")
                else:
                    area_trecho = area_do_arquivo(ids_validos[vid]["arquivo"])
                    if p["area"] != AREA_RELACIONAL and area_trecho not in (p["area"], AREA_RELACIONAL):
                        avisos.append(f"{pid}: trecho {vid[:8]} é da área '{area_trecho}', "
                                      f"mas a pergunta está marcada como '{p['area']}'")
    contagem = defaultdict(int)
    for p in perguntas:
        if p["pergunta"]:
            contagem[p["tipo"]] += 1
    print(f"Perguntas escritas: {sum(contagem.values())} "
          f"(literal {contagem['literal']}, parafrase {contagem['parafrase']}, sem_resposta {contagem['sem_resposta']})")
    for a in avisos:
        print("AVISO:", a)
    for e in erros:
        print("ERRO:", e)
    if erros:
        sys.exit(f"{len(erros)} erro(s) no gabarito.")
    print("Gabarito OK.")


# ----------------------------------------------------------------------------- métricas

def recall_at(ranking: list[str], relevantes: set[str], k: int) -> float:
    return len(set(ranking[:k]) & relevantes) / len(relevantes)


def rr_at(ranking: list[str], relevantes: set[str], k: int) -> float:
    for pos, vid in enumerate(ranking[:k], start=1):
        if vid in relevantes:
            return 1 / pos
    return 0.0


def percentil(valores: list[float], p: float) -> float:
    if not valores:
        return float("nan")
    ordenados = sorted(valores)
    if len(ordenados) == 1:
        return ordenados[0]
    pos = (len(ordenados) - 1) * p / 100
    baixo = int(pos)
    alto = min(baixo + 1, len(ordenados) - 1)
    return ordenados[baixo] + (ordenados[alto] - ordenados[baixo]) * (pos - baixo)


def fonte_da_pergunta(p: dict) -> str:
    return "relacional" if p["area"] == AREA_RELACIONAL else "pdf"


def escrever_csv(caminho: Path, linhas: list[dict]):
    caminho.parent.mkdir(parents=True, exist_ok=True)
    if not linhas:
        caminho.write_text("", encoding="utf-8")
        return
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()), delimiter=";")
        w.writeheader()
        for l in linhas:
            w.writerow({k: (f"{v:.4f}".replace(".", ",") if isinstance(v, float) else v) for k, v in l.items()})


def cmd_recuperar(args, api: Api):
    destino = Path(args.saida) if args.saida else RESULTADOS
    perguntas = [p for p in ler_gabarito(Path(args.gabarito)) if p["pergunta"]]
    com_resposta = [p for p in perguntas if p["tipo"] != "sem_resposta"]
    print(f"{len(perguntas)} perguntas ({len(com_resposta)} com resposta). Métodos: {', '.join(METODOS_RECUPERACAO)}; "
          f"{args.repeticoes} repetições para os tempos.")

    # Aquecimento: carrega o modelo de embeddings e as conexões antes de medir
    for m in METODOS_RECUPERACAO:
        api.recuperar(perguntas[0]["pergunta"], m, K_MAX)

    rankings, tempos = [], []
    ranking_por = {}
    for i, p in enumerate(perguntas, start=1):
        for m in METODOS_RECUPERACAO:
            for rep in range(args.repeticoes):
                r = api.recuperar(p["pergunta"], m, K_MAX)
                tempos.append({"pergunta_id": p["id"], "metodo": m, "repeticao": rep + 1,
                               "tempo_recuperacao_ms": r["tempoRecuperacaoMs"]})
                if rep == 0:
                    ids = [f["vectorId"] for f in r["fontes"]]
                    ranking_por[(p["id"], m)] = ids
                    for pos, f in enumerate(r["fontes"], start=1):
                        rankings.append({"pergunta_id": p["id"], "metodo": m, "posicao": pos,
                                         "vector_id": f["vectorId"], "arquivo": f["arquivo"],
                                         "chunk_index": f["chunkIndex"], "score": f["score"],
                                         "relevante": "sim" if f["vectorId"] in p["relevantes"] else "nao"})
        print(f"  [{i}/{len(perguntas)}] {p['id']}", flush=True)

    # Métricas por pergunta
    por_pergunta = []
    for p in com_resposta:
        rel = set(p["relevantes"])
        for m in METODOS_RECUPERACAO:
            rk = ranking_por[(p["id"], m)]
            linha = {"pergunta_id": p["id"], "tipo": p["tipo"], "area": p["area"], "fonte": fonte_da_pergunta(p),
                     "metodo": m}
            for k in KS:
                linha[f"recall@{k}"] = recall_at(rk, rel, k)
            linha["rr@10"] = rr_at(rk, rel, K_MAX)
            por_pergunta.append(linha)

    # Agregados: geral, por tipo e por fonte
    agregados = []
    grupos = [("geral", lambda l: "todas"), ("tipo", lambda l: l["tipo"]), ("fonte", lambda l: l["fonte"])]
    for nome_grupo, chave in grupos:
        valores = sorted({chave(l) for l in por_pergunta})
        for v in valores:
            for m in METODOS_RECUPERACAO:
                sel = [l for l in por_pergunta if l["metodo"] == m and chave(l) == v]
                if not sel:
                    continue
                linha = {"grupo": nome_grupo, "valor": v, "metodo": m, "n": len(sel)}
                for k in KS:
                    linha[f"recall@{k}"] = statistics.mean(l[f"recall@{k}"] for l in sel)
                linha["mrr@10"] = statistics.mean(l["rr@10"] for l in sel)
                agregados.append(linha)

    # Tempos de recuperação
    resumo_tempos = []
    for m in METODOS_RECUPERACAO:
        v = [t["tempo_recuperacao_ms"] for t in tempos if t["metodo"] == m]
        resumo_tempos.append({"metodo": m, "n": len(v), "mediana_ms": float(statistics.median(v)),
                              "p95_ms": float(percentil(v, 95)), "media_ms": float(statistics.mean(v))})

    # Concordância pgvector x ChromaDB (pergunta 2 da avaliação)
    concord = []
    for p in perguntas:
        a, b = ranking_por[(p["id"], "pgvector")], ranking_por[(p["id"], "chroma")]
        concord.append({"pergunta_id": p["id"], "mesma_ordem_top4": "sim" if a[:4] == b[:4] else "nao",
                        "mesma_ordem_top10": "sim" if a == b else "nao",
                        "intersecao_top10": len(set(a) & set(b))})
    n = len(concord)
    resumo_concord = {
        "perguntas": n,
        "mesma_ordem_top4": sum(c["mesma_ordem_top4"] == "sim" for c in concord),
        "mesma_ordem_top10": sum(c["mesma_ordem_top10"] == "sim" for c in concord),
        "intersecao_media_top10": statistics.mean(c["intersecao_top10"] for c in concord) if concord else 0,
    }

    escrever_csv(destino / "recuperacao_rankings.csv", rankings)
    escrever_csv(destino / "recuperacao_por_pergunta.csv", por_pergunta)
    escrever_csv(destino / "recuperacao_metricas.csv", agregados)
    escrever_csv(destino / "recuperacao_tempos.csv", tempos)
    escrever_csv(destino / "recuperacao_tempos_resumo.csv", resumo_tempos)
    escrever_csv(destino / "concordancia_pgvector_chroma.csv", concord)

    print("\nMétricas de recuperação (todas as perguntas com resposta):")
    print(f"  {'método':10} " + " ".join(f"{'R@' + str(k):>7}" for k in KS) + f" {'MRR@10':>7}")
    for l in agregados:
        if l["grupo"] == "geral":
            print(f"  {l['metodo']:10} " + " ".join(f"{l[f'recall@{k}']:7.3f}" for k in KS) + f" {l['mrr@10']:7.3f}")
    print("\nTempo de recuperação (ms):")
    for t in resumo_tempos:
        print(f"  {t['metodo']:10} mediana {t['mediana_ms']:8.1f}   p95 {t['p95_ms']:8.1f}")
    print(f"\npgvector x ChromaDB: mesmos 4 primeiros trechos na mesma ordem em {resumo_concord['mesma_ordem_top4']}"
          f"/{n} perguntas; mesmos 10 em {resumo_concord['mesma_ordem_top10']}/{n}.")
    print(f"\nArquivos em {destino}")


# ----------------------------------------------------------------------------- geração

ARQ_RESPOSTAS = "respostas.jsonl"
ARQ_CEGA = "avaliacao_cega.xlsx"
ARQ_CHAVE = "chave_avaliacao.csv"


def carregar_respostas(caminho: Path) -> dict:
    feitas = {}
    if caminho.exists():
        for linha in caminho.read_text(encoding="utf-8").splitlines():
            if linha.strip():
                r = json.loads(linha)
                feitas[(r["pergunta_id"], r["config"])] = r
    return feitas


def cmd_gerar(args, api: Api):
    perguntas = [p for p in ler_gabarito(Path(args.gabarito)) if p["pergunta"]]
    RESULTADOS.mkdir(parents=True, exist_ok=True)
    arq = RESULTADOS / ARQ_RESPOSTAS
    feitas = carregar_respostas(arq)
    tarefas = []
    for p in perguntas:
        for c in CONFIGS_GERACAO:
            if c == "teto" and p["tipo"] == "sem_resposta":
                continue
            if (p["id"], c) not in feitas:
                tarefas.append((p, c))
    print(f"{len(feitas)} respostas já geradas; faltam {len(tarefas)}. "
          f"Pode interromper (Ctrl+C) e rodar de novo: o script continua de onde parou.")

    if tarefas:
        # Aquecimento: carrega o modelo de linguagem na memória antes de medir
        print("Aquecendo o modelo...", flush=True)
        api.consultar(perguntas[0]["pergunta"], "pgvector")

    with arq.open("a", encoding="utf-8") as saida:
        for i, (p, c) in enumerate(tarefas, start=1):
            inicio = time.time()
            if c == "teto":
                r = api.teto(p["pergunta"], p["relevantes"])
            else:
                r = api.consultar(p["pergunta"], c)
            reg = {"pergunta_id": p["id"], "config": c, "resposta": r["resposta"],
                   "fontes": [{"vectorId": f["vectorId"], "arquivo": f["arquivo"], "chunkIndex": f["chunkIndex"],
                               "trecho": f["trecho"]} for f in r["fontes"]],
                   "tempo_recuperacao_ms": r["tempoRecuperacaoMs"], "tempo_geracao_ms": r["tempoGeracaoMs"]}
            saida.write(json.dumps(reg, ensure_ascii=False) + "\n")
            saida.flush()
            feitas[(p["id"], c)] = reg
            print(f"  [{i}/{len(tarefas)}] {p['id']} {c:8} {time.time() - inicio:6.1f}s", flush=True)

    montar_planilha_cega(perguntas, feitas, args.semente)


def montar_planilha_cega(perguntas: list[dict], feitas: dict, semente: int):
    destino = RESULTADOS / ARQ_CEGA
    if destino.exists():
        print(f"{destino} já existe e não foi recriado (para não apagar notas já dadas). "
              f"Apague o arquivo se quiser gerá-lo de novo.")
        return
    itens = {}
    for p in perguntas:
        for c in CONFIGS_GERACAO:
            r = feitas.get((p["id"], c))
            if r is None:
                continue
            # Mesma resposta com o mesmo contexto => avaliada uma vez só
            chave = (p["id"], r["resposta"].strip(), tuple(f["vectorId"] for f in r["fontes"]))
            if chave not in itens:
                itens[chave] = {"pergunta": p, "registro": r, "configs": []}
            itens[chave]["configs"].append(c)
    lista = list(itens.values())
    random.Random(semente).shuffle(lista)

    wb = Workbook()
    ws = wb.active
    ws.title = "Avaliacao"
    cabecalho(ws, [("id_cego", 9), ("pergunta", 45), ("sem_resposta_na_base", 12), ("resposta_referencia", 40),
                   ("resposta_gerada", 60), ("trechos_enviados_ao_modelo", 80), ("nota", 7), ("fundamentada", 13),
                   ("informou_ausencia", 13), ("comentario", 30)])
    chave_linhas = []
    for i, item in enumerate(lista, start=1):
        p, r = item["pergunta"], item["registro"]
        idc = f"R{i:03d}"
        trechos = "\n\n".join(f"[{j}] {f['arquivo']} (trecho {f['chunkIndex']})\n{f['trecho']}"
                              for j, f in enumerate(r["fontes"], start=1)) or "(nenhum trecho)"
        sem = p["tipo"] == "sem_resposta"
        ws.append([idc, p["pergunta"], "sim" if sem else "nao", p["referencia"] or "-", r["resposta"],
                   trechos[:32000], None, None, None, None])
        for c in ws[ws.max_row]:
            c.alignment = QUEBRA
        ws.row_dimensions[ws.max_row].height = 150
        chave_linhas.append({"id_cego": idc, "pergunta_id": p["id"], "configs": "|".join(item["configs"])})
    n = len(lista) + 1
    dv_nota = DataValidation(type="list", formula1='"' + ",".join(map(str, NOTAS_VALIDAS)) + '"')
    dv_sn = DataValidation(type="list", formula1='"sim,nao"')
    ws.add_data_validation(dv_nota)
    ws.add_data_validation(dv_sn)
    dv_nota.add(f"G2:G{n}")
    dv_sn.add(f"H2:I{n}")

    leg = wb.create_sheet("Escala")
    leg.append(["nota", "critério (Finardi et al., 2024)"])
    for nota, crit in [(1, "A resposta não tem relação com a referência."),
                       (3, "A resposta tem pouca relevância e não está alinhada com a referência."),
                       (5, "A resposta tem relevância moderada, mas contém imprecisões."),
                       (7, "A resposta está alinhada com a referência, mas tem pequenas omissões."),
                       (10, "A resposta está totalmente correta e alinhada com a referência.")]:
        leg.append([nota, crit])
    leg.append([])
    leg.append(["fundamentada", "sim: todas as afirmações da resposta estão nos trechos enviados; nao: alguma afirmação não está."])
    leg.append(["informou_ausencia", "só para perguntas sem resposta na base: o modelo disse que não encontrou a informação?"])
    leg.append(["nota (sem resposta)", "nas perguntas sem resposta, deixe a nota vazia; avalie só 'informou_ausencia' e 'fundamentada'."])
    leg.column_dimensions["A"].width = 22
    leg.column_dimensions["B"].width = 110
    for c in leg[1]:
        c.font = Font(bold=True)

    wb.save(destino)
    escrever_csv(RESULTADOS / ARQ_CHAVE, chave_linhas)
    total = sum(len(i["configs"]) for i in lista)
    print(f"\nPlanilha de avaliação cega: {destino}")
    print(f"  {len(lista)} respostas para avaliar ({total - len(lista)} repetidas foram agrupadas).")
    print(f"  Não abra {ARQ_CHAVE} antes de terminar as notas: ele diz qual método gerou cada resposta.")


def cmd_notas(args, _api=None):
    perguntas = {p["id"]: p for p in ler_gabarito(Path(args.gabarito)) if p["pergunta"]}
    ws = load_workbook(RESULTADOS / ARQ_CEGA, data_only=True)["Avaliacao"]
    nomes = [c.value for c in ws[1]]
    avaliacoes = {}
    pendentes = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        reg = dict(zip(nomes, linha))
        if not reg.get("id_cego"):
            continue
        sem = reg.get("sem_resposta_na_base") == "sim"
        nota = reg.get("nota")
        fund = (reg.get("fundamentada") or "").strip().lower()
        aus = (reg.get("informou_ausencia") or "").strip().lower()
        if (not sem and nota not in NOTAS_VALIDAS) or fund not in ("sim", "nao") or (sem and aus not in ("sim", "nao")):
            pendentes.append(reg["id_cego"])
        avaliacoes[reg["id_cego"]] = {"nota": nota, "fundamentada": fund, "informou_ausencia": aus}
    if pendentes:
        print(f"AVISO: {len(pendentes)} resposta(s) sem avaliação completa: {', '.join(pendentes[:15])}"
              f"{'...' if len(pendentes) > 15 else ''}")

    with (RESULTADOS / ARQ_CHAVE).open(encoding="utf-8-sig") as f:
        chave = list(csv.DictReader(f, delimiter=";"))
    por_config = defaultdict(list)  # config -> [(pergunta, avaliacao)]
    for c in chave:
        av = avaliacoes.get(c["id_cego"])
        if av is None:
            continue
        for cfg in c["configs"].split("|"):
            por_config[cfg].append((perguntas[c["pergunta_id"]], av))

    respostas = carregar_respostas(RESULTADOS / ARQ_RESPOSTAS)

    def resumo(cfg: str, filtro, rotulo: str) -> dict:
        itens = [(p, a) for p, a in por_config.get(cfg, []) if filtro(p)]
        com = [a for p, a in itens if p["tipo"] != "sem_resposta" and a["nota"] in NOTAS_VALIDAS]
        sem = [a for p, a in itens if p["tipo"] == "sem_resposta" and a["informou_ausencia"] in ("sim", "nao")]
        fund = [a for _, a in itens if a["fundamentada"] in ("sim", "nao")]
        return {"config": cfg, "grupo": rotulo,
                "n_com_resposta": len(com),
                "nota_media": statistics.mean(a["nota"] for a in com) if com else float("nan"),
                "taxa_fundamentadas": (sum(a["fundamentada"] == "sim" for a in fund) / len(fund)) if fund else float("nan"),
                "taxa_alucinacao": (sum(a["fundamentada"] == "nao" for a in fund) / len(fund)) if fund else float("nan"),
                "n_sem_resposta": len(sem),
                "taxa_informou_ausencia": (sum(a["informou_ausencia"] == "sim" for a in sem) / len(sem)) if sem else float("nan")}

    linhas = []
    filtros = [("geral", lambda p: True), ("literal", lambda p: p["tipo"] == "literal"),
               ("parafrase", lambda p: p["tipo"] == "parafrase"),
               ("relacional", lambda p: p["area"] == AREA_RELACIONAL),
               ("pdf", lambda p: p["tipo"] != "sem_resposta" and p["area"] != AREA_RELACIONAL)]
    for rotulo, filtro in filtros:
        teto = resumo("teto", filtro, rotulo)["nota_media"]
        for cfg in CONFIGS_GERACAO:
            l = resumo(cfg, filtro, rotulo)
            l["degradacao_vs_teto"] = (1 - l["nota_media"] / teto) if cfg != "teto" and teto == teto and teto else float("nan")
            linhas.append(l)

    tempos = []
    for cfg in CONFIGS_GERACAO:
        g = [r["tempo_geracao_ms"] for (pid, c), r in respostas.items() if c == cfg]
        rec = [r["tempo_recuperacao_ms"] for (pid, c), r in respostas.items() if c == cfg]
        if g:
            tempos.append({"config": cfg, "n": len(g),
                           "geracao_mediana_ms": float(statistics.median(g)), "geracao_p95_ms": float(percentil(g, 95)),
                           "recuperacao_mediana_ms": float(statistics.median(rec)),
                           "participacao_geracao_no_total": sum(g) / (sum(g) + sum(rec)) if (sum(g) + sum(rec)) else float("nan")})

    escrever_csv(RESULTADOS / "respostas_metricas.csv", linhas)
    escrever_csv(RESULTADOS / "geracao_tempos_resumo.csv", tempos)

    print(f"\n{'config':9} {'grupo':10} {'nota':>6} {'degrad.':>8} {'fund.':>6} {'aluc.':>6} {'aus.':>6}")
    for l in linhas:
        def f(v, pct=False):
            if v != v:
                return "     -"
            return f"{v * 100:5.1f}%" if pct else f"{v:6.2f}"
        print(f"{l['config']:9} {l['grupo']:10} {f(l['nota_media'])} {f(l['degradacao_vs_teto'], True):>8} "
              f"{f(l['taxa_fundamentadas'], True)} {f(l['taxa_alucinacao'], True)} {f(l['taxa_informou_ausencia'], True)}")
    print("\nTempo de geração (ms):")
    for t in tempos:
        print(f"  {t['config']:9} mediana {t['geracao_mediana_ms']:9.0f}   p95 {t['geracao_p95_ms']:9.0f}   "
              f"geração = {t['participacao_geracao_no_total'] * 100:.1f}% do tempo total")
    print(f"\nArquivos em {RESULTADOS}")


# ----------------------------------------------------------------------------- CLI

def main():
    ap = argparse.ArgumentParser(description="Experimento do TCC: busca semântica e RAG (pgvector x ChromaDB x lexical)")
    ap.add_argument("--url", default="http://localhost:8080", help="endereço da aplicação (padrão: %(default)s)")
    ap.add_argument("--timeout", type=float, default=600, help="tempo máximo por requisição, em segundos")
    ap.add_argument("--gabarito", default=str(AQUI / "gabarito.xlsx"))
    sub = ap.add_subparsers(dest="comando", required=True)

    s = sub.add_parser("trechos", help="exporta todos os trechos da base para trechos.xlsx")
    s.add_argument("--saida", default=str(AQUI / "trechos.xlsx"))
    s = sub.add_parser("modelo", help="cria a planilha do gabarito (gabarito.xlsx)")
    s.add_argument("--saida", default=str(AQUI / "gabarito.xlsx"))
    s.add_argument("--sobrescrever", action="store_true")
    s = sub.add_parser("validar", help="confere o gabarito")
    s.add_argument("--offline", action="store_true", help="não consulta a aplicação para conferir os vector_id")
    s = sub.add_parser("recuperar", help="métricas e tempos de recuperação dos três métodos")
    s.add_argument("--repeticoes", type=int, default=5, help="repetições de cada consulta para medir o tempo")
    s.add_argument("--saida", default=None, help="pasta dos resultados (padrão: resultados/)")
    s = sub.add_parser("gerar", help="gera as respostas (pgvector, chroma e teto) e a planilha de avaliação cega")
    s.add_argument("--semente", type=int, default=2026, help="semente do embaralhamento da planilha")
    sub.add_parser("notas", help="calcula as métricas de resposta a partir das notas")

    args = ap.parse_args()
    api = Api(args.url, args.timeout)
    comandos = {"trechos": cmd_trechos, "modelo": cmd_modelo, "validar": cmd_validar,
                "recuperar": cmd_recuperar, "gerar": cmd_gerar, "notas": cmd_notas}
    try:
        comandos[args.comando](args, api)
    except RuntimeError as e:
        sys.exit(f"Erro: {e}")


if __name__ == "__main__":
    main()
